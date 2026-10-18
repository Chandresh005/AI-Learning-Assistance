import os
from typing import List, Dict, Any
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.chunk import DocumentChunk
from app.models.resource import Resource


def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text page-by-page from a PDF file using PyMuPDF.
    Returns list of {'page_number': int, 'text': str}.
    """
    import pymupdf

    doc = pymupdf.open(file_path)
    pages = []
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        text = page.get_text().strip()
        if text:
            pages.append({
                "page_number": page_idx + 1,  # 1-indexed
                "text": text
            })
    doc.close()
    return pages


def extract_text_from_pptx(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text slide-by-slide from a PPTX file using python-pptx.
    Returns list of {'slide_number': int, 'text': str}.
    """
    from pptx import Presentation

    prs = Presentation(file_path)
    slides = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        full_slide_text = "\n".join(slide_texts).strip()
        if full_slide_text:
            slides.append({
                "slide_number": slide_idx + 1,  # 1-indexed
                "text": full_slide_text
            })
    return slides


def chunk_document_elements(
    elements: List[Dict[str, Any]],
    source_type: str,
    max_chunk_chars: int = 800
) -> List[Dict[str, Any]]:
    """
    Chunks document elements (pages or slides) into discrete retrieval chunks,
    preserving exact page_number or slide_number metadata on every chunk.
    """
    chunks = []
    chunk_index = 1

    for elem in elements:
        text = elem.get("text", "").strip()
        page_num = elem.get("page_number")
        slide_num = elem.get("slide_number")

        if not text:
            continue

        # If page or slide is shorter than max_chunk_chars, keep as one chunk
        if len(text) <= max_chunk_chars:
            chunks.append({
                "chunk_id": chunk_index,
                "text": text,
                "source_type": source_type,
                "page_number": page_num,
                "slide_number": slide_num
            })
            chunk_index += 1
        else:
            # Split longer text into paragraphs or blocks
            paragraphs = text.split("\n\n")
            curr_text = ""
            for p in paragraphs:
                p = p.strip()
                if not p:
                    continue
                if len(curr_text) + len(p) <= max_chunk_chars:
                    curr_text = f"{curr_text}\n\n{p}".strip()
                else:
                    if curr_text:
                        chunks.append({
                            "chunk_id": chunk_index,
                            "text": curr_text,
                            "source_type": source_type,
                            "page_number": page_num,
                            "slide_number": slide_num
                        })
                        chunk_index += 1
                    curr_text = p

            if curr_text:
                chunks.append({
                    "chunk_id": chunk_index,
                    "text": curr_text,
                    "source_type": source_type,
                    "page_number": page_num,
                    "slide_number": slide_num
                })
                chunk_index += 1

    return chunks


async def store_document_chunks(
    course_id: str,
    resource_id: str,
    chunks: List[Dict[str, Any]],
    db: AsyncSession
) -> int:
    """
    Persists document chunks with page/slide metadata into document_chunks table.
    """
    for c in chunks:
        record = DocumentChunk(
            course_id=course_id,
            resource_id=resource_id,
            chunk_id=c["chunk_id"],
            text=c["text"],
            source_type=c["source_type"],
            page_number=c.get("page_number"),
            slide_number=c.get("slide_number")
        )
        db.add(record)

    await db.commit()
    return len(chunks)
